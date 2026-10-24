import re
import datetime
import difflib
import openpyxl
from django.utils import timezone
from openpyxl.utils.datetime import from_excel

from employees.models import Employee


def guess_counselor_name(sheet_name):
    """Strip common prefixes/suffixes from a sheet name to get a likely counselor first name."""
    name = sheet_name
    name = re.sub(r"(?i)^(sir|mam|miss|mr|mrs)\.?\s+", "", name.strip())
    name = re.sub(r"(?i)('?s)?\s+students?$", "", name).strip()
    return name


def _words(name):
    """Name words in lower case, without the 'Demo -' prefix of demo data or punctuation."""
    name = re.sub(r'(?i)^demo\s*-\s*', '', name or '')
    return re.findall(r'[a-z]+', name.lower())


def match_employee(name_hint, employees_qs):
    """
    Match a sheet name or column heading to an employee. Returns (employee_or_none, confidence).
    Whole words only: "Ali" matches "Ali Raza" but never "Abdul Malik". If two employees match
    equally well (two people called Ali), nobody is picked and the admin chooses.
    """
    hint = _words(name_hint)
    if not hint:
        return None, 0
    employees = list(employees_qs)

    exact = [e for e in employees if _words(e.name) == hint]
    if len(exact) == 1:
        return exact[0], 'high'
    if len(exact) > 1:
        return None, 0

    contains = [e for e in employees if all(word in _words(e.name) for word in hint)]
    if len(contains) == 1:
        return contains[0], 'high'
    if len(contains) > 1:
        return None, 0

    # Spelling differences ("Hinna" for "Hina"): compare with full names and with first names.
    candidates = {}
    for e in employees:
        words = _words(e.name)
        if words:
            candidates.setdefault(' '.join(words), []).append(e)
            candidates.setdefault(words[0], []).append(e)
    close = difflib.get_close_matches(' '.join(hint), list(candidates), n=1, cutoff=0.75)
    if close and len(candidates[close[0]]) == 1:
        return candidates[close[0]][0], 'fuzzy'
    return None, 0


DATE_FORMATS = ('%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y-%m-%d', '%d/%m/%y', '%d %b %Y', '%d %B %Y', '%d-%b-%Y', '%d-%b-%y')


def parse_date_cell(value):
    """Real date cells, Excel serial numbers, and common text formats (day first, as used in Pakistan)."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial date (1954-2119)
        return (from_excel(value)).date()
    if isinstance(value, str):
        value = ' '.join(value.strip().split())
        for fmt in DATE_FORMATS:
            try:
                return datetime.datetime.strptime(value, fmt).date()
            except ValueError:
                continue
    return None


def parse_time_cell(value):
    """A check-in time from a time cell, a date-time cell, a fraction of a day, or text like '10:05' / '10:05 AM'."""
    if isinstance(value, datetime.datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, datetime.time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, float) and 0 < value < 1:
        minutes = round(value * 24 * 60)
        return datetime.time(minutes // 60, minutes % 60)
    if isinstance(value, str):
        m = re.match(r'^\s*(\d{1,2})[:.](\d{2})\s*([ap]\.?m\.?)?', value, re.I)
        if m:
            hour, minute = int(m.group(1)), int(m.group(2))
            suffix = (m.group(3) or '').lower().replace('.', '')
            if suffix == 'pm' and hour < 12:
                hour += 12
            if suffix == 'am' and hour == 12:
                hour = 0
            if hour < 24 and minute < 60:
                return datetime.time(hour, minute)
    return None


def parse_enrollment_workbook(file_obj):
    """
    Parses a workbook where each sheet is one counselor's students,
    with columns: Name, Enrollment Date, Country.
    Returns a list of sheet result dicts for preview.
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    active_employees = Employee.objects.filter(is_active=True)
    results = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header = [str(c).strip().lower() if c else '' for c in rows[0]]
        try:
            name_idx = next(i for i, h in enumerate(header) if 'name' in h)
            date_idx = next(i for i, h in enumerate(header) if 'date' in h)
            country_idx = next(i for i, h in enumerate(header) if 'country' in h)
        except StopIteration:
            results.append({
                'sheet_name': sheet_name, 'counselor_hint': None, 'matched_employee': None,
                'match_confidence': 0, 'rows': [], 'error': 'Could not find Name/Enrollment Date/Country columns in this sheet.'
            })
            continue

        counselor_hint = guess_counselor_name(sheet_name)
        matched_employee, confidence = match_employee(counselor_hint, active_employees)

        today = timezone.localdate()
        parsed_rows = []
        for row in rows[1:]:
            if not row or not row[name_idx] or not str(row[name_idx]).strip():
                continue
            student_name = ' '.join(str(row[name_idx]).split())
            country = ' '.join(str(row[country_idx]).split()) if row[country_idx] else ''
            confirmed_on = parse_date_cell(row[date_idx])
            problem = ('date not recognised' if confirmed_on is None else
                       'date is in the future' if confirmed_on > today else
                       'country missing' if not country else '')
            parsed_rows.append({
                'student_name': student_name,
                'country': country,
                'confirmed_on': confirmed_on,
                'date_error': confirmed_on is None,
                'problem': problem,
            })

        results.append({
            'sheet_name': sheet_name,
            'counselor_hint': counselor_hint,
            'matched_employee': matched_employee,
            'match_confidence': confidence,
            'rows': parsed_rows,
            'error': None,
        })

    return results


def parse_attendance_workbook(file_obj):
    """
    Parses the ALHAM attendance matrix format:
    - A header block containing 'Year' / 'Month:' somewhere near the top
    - A row with employee column numbers, then a row with employee names
    - Data rows: day number, day name, then one cell per employee
    - 'SUNDAY' (or similar) merged rows mean holiday, no per-employee cells
    - Footer rows: Present / Absent / Holiday totals (ignored for import)
    """
    import re as _re
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    month_num, year_num = None, None
    header_row_idx = None
    for i, row in enumerate(rows[:6]):
        for cell in row:
            if cell and isinstance(cell, str) and 'month' in cell.lower():
                m = _re.search(r'([A-Za-z]+)\s+(\d{4})', cell)
                if m:
                    for fmt in ('%B', '%b'):  # "October" or "Oct"
                        try:
                            month_num = datetime.datetime.strptime(m.group(1), fmt).month
                            break
                        except ValueError:
                            continue
                    year_num = int(m.group(2))
        for cell in row:
            if cell and isinstance(cell, str) and 'date' in cell.lower():
                header_row_idx = i
                break
        if header_row_idx is not None:
            break

    if header_row_idx is None or month_num is None:
        return {'error': 'Could not detect the month/year or header row in this sheet. Please check the file format.'}

    # Each employee keeps the column their name is in, so a blank column between names can't shift data.
    name_row = rows[header_row_idx + 1]
    name_columns = [(idx, str(c).strip()) for idx, c in enumerate(name_row) if idx >= 2 and c and str(c).strip()]
    employee_names = [name for _, name in name_columns]

    active_employees = Employee.objects.filter(is_active=True)
    employee_map = {name: match_employee(name, active_employees)[0] for name in employee_names}
    unmatched_names = {name for name, emp in employee_map.items() if emp is None}

    today = timezone.localdate()
    entries = []
    skipped_future = 0
    for row in rows[header_row_idx + 2:]:
        if not row or not isinstance(row[0], (int, float)) or isinstance(row[0], bool):
            continue
        try:
            the_date = datetime.date(year_num, month_num, int(row[0]))
        except ValueError:
            continue
        if the_date > today:
            skipped_future += 1
            continue
        day_str = str(row[1]).strip().upper() if len(row) > 1 and row[1] else ''
        cells = [row[idx] if idx < len(row) else None for idx, _ in name_columns]
        row_text = ' '.join(str(c) for c in row[1:] if isinstance(c, str)).lower()

        whole_row_holiday = 'holiday' in row_text and all(c is None or isinstance(c, str) for c in cells)
        if 'SUN' in day_str or the_date.weekday() == 6 or whole_row_holiday:
            for name in employee_names:
                entries.append({'employee_name': name, 'employee': employee_map.get(name), 'date': the_date,
                                'status': 'holiday', 'check_in_time': None, 'note': '', 'raw': None})
            continue
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in cells):
            continue  # a day nobody filled in yet: leave it open rather than guessing

        for (_, name), value in zip(name_columns, cells):
            entry = {'employee_name': name, 'employee': employee_map.get(name), 'date': the_date,
                     'raw': value, 'note': '', 'check_in_time': None, 'status': 'needs_review'}
            text = value.strip().lower() if isinstance(value, str) else ''
            check_in = parse_time_cell(value)
            if value is None or (isinstance(value, str) and not text):
                entry['note'] = 'blank in sheet'
            elif 'not mention' in text:
                entry['note'] = 'not mentioned in sheet'
            elif 'leave' in text:
                entry['status'] = 'leave'
            elif 'holiday' in text or text == 'off':
                entry['status'] = 'holiday'
            elif 'absent' in text or text == 'a':
                entry['status'] = 'absent'
            elif check_in is not None:
                entry['status'] = 'present'
                entry['check_in_time'] = check_in
                if isinstance(value, str) and not re.fullmatch(r'\s*\d{1,2}[:.]\d{2}\s*', value):
                    entry['note'] = value.strip()[:100]  # e.g. "10:40 (late)" keeps its remark
            else:
                entry['note'] = str(value)[:100]
            entries.append(entry)

    return {
        'error': None,
        'month': month_num,
        'year': year_num,
        'employee_names': employee_names,
        'unmatched_names': sorted(unmatched_names),
        'entries': entries,
        'skipped_future': skipped_future,
    }