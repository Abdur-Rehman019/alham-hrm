import datetime
import io
import re
from decimal import Decimal
from unittest import mock

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from attendance.models import DailyAttendance, Leave
from employees.models import Employee
from payroll.models import IncentiveRule, StudentEnrollment

PW = 'Str0ng!Pass#2026'
TODAY = datetime.date(2026, 10, 24)  # a Saturday; imports are checked as of this day


def xlsx(build):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile('data.xlsx', buf.getvalue(),
                              content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def make_employee(name, **kw):
    return Employee.objects.create(name=name, role=kw.pop('role', 'counselor'), monthly_salary=30000,
                                   joined_on=datetime.date(2024, 1, 1), **kw)


class ImportBase(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser('boss', 'boss@alham.pk', PW)
        self.client.force_login(self.admin)
        patcher = mock.patch('django.utils.timezone.localdate', return_value=TODAY)
        patcher.start()
        self.addCleanup(patcher.stop)


class EnrollmentImportTests(ImportBase):
    def setUp(self):
        super().setUp()
        self.hina = make_employee('Hina Malik')
        self.ali = make_employee('Ali Raza')
        IncentiveRule.objects.create(destination_country='UK', amount=Decimal('10000'))

    def _workbook(self, sheets):
        def build(wb):
            for title, rows in sheets.items():
                ws = wb.create_sheet(title)
                ws.append(['Student Name', 'Enrollment Date', 'Country'])
                for row in rows:
                    ws.append(row)
        return xlsx(build)

    def _preview(self, upload):
        response = self.client.post('/imports/enrollments/', {'file': upload})
        self.assertEqual(response.status_code, 200)
        return response

    def _confirm(self, response, **overrides):
        token = re.search(r'name="token" value="([0-9a-f]{32})"', response.content.decode()).group(1)
        data = {'token': token}
        for i, sheet in enumerate(response.context['sheets']):
            if sheet['matched_employee']:
                data[f'counselor_{i}'] = sheet['matched_employee'].pk
        data.update(overrides)
        return self.client.post('/imports/enrollments/confirm/', data, follow=True)

    def test_happy_path_assigns_incentive(self):
        upload = self._workbook({'Mam Hina Students': [['Sara', datetime.datetime(2026, 10, 3), 'UK']]})
        response = self._preview(upload)
        self.assertEqual(response.context['sheets'][0]['matched_employee'], self.hina)
        self._confirm(response)
        enrollment = StudentEnrollment.objects.get()
        self.assertEqual(enrollment.counselor, self.hina)
        self.assertEqual(enrollment.incentive_amount, Decimal('10000'))

    def test_sheet_name_does_not_match_inside_another_name(self):
        # "Ali" is inside "mALIk" and "Abdul Malik" sorts first; the sheet must still go to Ali Raza.
        make_employee('Abdul Malik')
        response = self._preview(self._workbook({'Sir Ali Students': [['Sara', '03/10/2026', 'UK']]}))
        self.assertEqual(response.context['sheets'][0]['matched_employee'], self.ali)

    def test_importing_the_same_file_twice_does_not_duplicate(self):
        rows = {'Hina': [['Sara', '03/10/2026', 'UK'], ['Omar', '04/10/2026', 'UK']]}
        self._confirm(self._preview(self._workbook(rows)))
        response = self._confirm(self._preview(self._workbook(rows)))
        self.assertEqual(StudentEnrollment.objects.count(), 2)
        self.assertContains(response, 'already imported')

    def test_common_date_formats(self):
        rows = {'Hina': [
            ['A', '03/10/2026', 'UK'], ['B', '3-10-2026', 'UK'], ['C', '2026-10-03', 'UK'],
            ['D', '03.10.2026', 'UK'], ['E', '3 Oct 2026', 'UK'], ['F', 46298, 'UK'],  # Excel serial number
        ]}
        response = self._preview(self._workbook(rows))
        dates = [r['confirmed_on'] for r in response.context['sheets'][0]['rows']]
        self.assertEqual(dates, [datetime.date(2026, 10, 3)] * 5 + [datetime.date(2026, 10, 3)])

    def test_future_and_bad_rows_are_reported_not_silently_dropped(self):
        rows = {'Hina': [['A', '03/10/2026', 'UK'], ['B', '30/12/2026', 'UK'], ['C', 'soon', 'UK'], ['D', '03/10/2026', '']]}
        response = self._confirm(self._preview(self._workbook(rows)))
        self.assertEqual(StudentEnrollment.objects.count(), 1)
        self.assertContains(response, 'Skipped 3')

    def test_expired_token_and_garbage_token(self):
        response = self.client.post('/imports/enrollments/confirm/', {'token': '../../etc/passwd'}, follow=True)
        self.assertContains(response, 'Upload session expired')
        response = self.client.post('/imports/enrollments/confirm/', {'token': 'a' * 32}, follow=True)
        self.assertContains(response, 'Upload session expired')

    def test_sheet_without_expected_columns_shows_error(self):
        def build(wb):
            ws = wb.create_sheet('Hina')
            ws.append(['Student', 'When', 'Where'])
            ws.append(['Sara', '03/10/2026', 'UK'])
        response = self._preview(xlsx(build))
        self.assertContains(response, 'Could not find')


class AttendanceImportTests(ImportBase):
    def setUp(self):
        super().setUp()
        self.hina = make_employee('Hina Malik')
        self.ali = make_employee('Ali Raza')

    def _sheet(self, names, days, name_columns=None):
        """ALHAM matrix: title rows, 'Date' header row, a names row, then one row per day."""
        def build(wb):
            ws = wb.create_sheet('Attendance')
            ws.append(['ALHAM Overseas Consultants'])
            ws.append(['Month: October 2026'])
            ws.append(['Date', 'Day'])
            cols = name_columns or list(range(2, 2 + len(names)))
            name_row = [None] * (max(cols) + 1)
            for col, name in zip(cols, names):
                name_row[col] = name
            ws.append(name_row)
            for day, day_name, values in days:
                row = [day, day_name] + [None] * (max(cols) - 1)
                for col, value in zip(cols, values):
                    row[col] = value
                ws.append(row)
        return xlsx(build)

    def _import(self, upload):
        return self.client.post('/imports/attendance/', {'file': upload}, follow=True)

    def record(self, employee, day):
        return DailyAttendance.objects.get(employee=employee, date=datetime.date(2026, 10, day))

    def test_basic_statuses(self):
        self._import(self._sheet(['Hina', 'Ali'], [
            (1, 'THURSDAY', [datetime.time(10, 5), 'Leave']),
            (2, 'FRIDAY', ['10:40', 'not mention']),
            (4, 'SUNDAY', [None, None]),
        ]))
        self.assertEqual(self.record(self.hina, 1).status, 'present')
        self.assertFalse(self.record(self.hina, 1).is_late)
        self.assertEqual(self.record(self.ali, 1).status, 'leave')
        self.assertTrue(Leave.objects.filter(employee=self.ali).exists())
        self.assertTrue(self.record(self.hina, 2).is_late)
        self.assertEqual(self.record(self.ali, 2).status, 'needs_review')
        self.assertEqual(self.record(self.hina, 4).status, 'holiday')

    def test_time_stored_as_datetime_or_fraction(self):
        self._import(self._sheet(['Hina', 'Ali'], [
            (1, 'THURSDAY', [datetime.datetime(1900, 1, 1, 10, 5), 0.4375]),  # 0.4375 of a day = 10:30
        ]))
        self.assertEqual(self.record(self.hina, 1).check_in_time, datetime.time(10, 5))
        self.assertEqual(self.record(self.ali, 1).check_in_time, datetime.time(10, 30))
        self.assertEqual(self.record(self.ali, 1).status, 'present')

    def test_blank_weekday_row_is_not_turned_into_a_holiday(self):
        # Days not filled in yet (e.g. later this month) must stay open, not become holidays.
        self._import(self._sheet(['Hina', 'Ali'], [
            (1, 'THURSDAY', [datetime.time(10, 0), datetime.time(10, 0)]),
            (20, 'TUESDAY', [None, None]),
        ]))
        self.assertFalse(DailyAttendance.objects.filter(date=datetime.date(2026, 10, 20)).exists())

    def test_future_days_are_not_imported(self):
        self._import(self._sheet(['Hina'], [(28, 'WEDNESDAY', [datetime.time(10, 0)])]))
        self.assertFalse(DailyAttendance.objects.exists())

    def test_gap_between_name_columns_keeps_values_aligned(self):
        self._import(self._sheet(['Hina', 'Ali'], [(1, 'THURSDAY', [datetime.time(10, 5), 'Leave'])],
                                 name_columns=[2, 4]))
        self.assertEqual(self.record(self.hina, 1).status, 'present')
        self.assertEqual(self.record(self.ali, 1).status, 'leave')

    def test_name_inside_another_name_is_not_matched(self):
        abdul = make_employee('Abdul Malik')
        self._import(self._sheet(['Ali'], [(1, 'THURSDAY', [datetime.time(10, 0)])]))
        self.assertTrue(DailyAttendance.objects.filter(employee=self.ali).exists())
        self.assertFalse(DailyAttendance.objects.filter(employee=abdul).exists())

    def test_reimport_reports_existing_rows(self):
        upload = lambda: self._sheet(['Hina'], [(1, 'THURSDAY', [datetime.time(10, 0)])])
        self._import(upload())
        response = self._import(upload())
        self.assertEqual(DailyAttendance.objects.count(), 1)
        self.assertContains(response, 'already recorded')

    def test_unreadable_sheet_shows_error(self):
        def build(wb):
            wb.create_sheet('Sheet').append(['nothing useful'])
        response = self._import(xlsx(build))
        self.assertContains(response, 'Could not detect the month')
